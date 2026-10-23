"""In-memory parsing of small local resume files; never execute document content."""

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader

from .errors import InputError
from .validation import MAX_RESUME_CHARS, validate_text

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_BYTES = 20 * 1024 * 1024
MAX_PAGES = 20


def extract_resume(data: bytes, filename: str) -> str:
    if not data:
        raise InputError("The uploaded file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise InputError("Upload a file smaller than 5 MB.")
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".pdf":
            if not data.startswith(b"%PDF-"):
                raise InputError("This file is not a valid PDF. Export it again or paste the text.")
            reader = PdfReader(BytesIO(data), strict=True)
            if reader.is_encrypted:
                raise InputError("Password-protected PDFs are unsupported. Upload an unlocked copy.")
            if len(reader.pages) > MAX_PAGES:
                raise InputError("Use a resume PDF with 20 pages or fewer.")
            parts = []
            for page in reader.pages:
                text = page.extract_text() or ""
                parts.append(text)
                if sum(map(len, parts)) > MAX_RESUME_CHARS:
                    raise InputError("The extracted resume is too long. Use a shorter file.")
            text = "\n".join(parts)
        elif suffix == ".docx":
            if not data.startswith(b"PK"):
                raise InputError("This file is not a valid DOCX. Export it again or paste the text.")
            with ZipFile(BytesIO(data)) as archive:
                info = archive.infolist()
                if len(info) > 200 or sum(x.file_size for x in info) > MAX_EXPANDED_BYTES or any(x.flag_bits & 1 for x in info):
                    raise InputError("This DOCX is too complex, compressed, or encrypted. Paste the text instead.")
                if "word/document.xml" not in archive.namelist():
                    raise InputError("This archive is not a Word DOCX document.")
                if any("vbaproject" in x.filename.lower() for x in info):
                    raise InputError("Macro-containing documents are unsupported. Paste the text instead.")
            document = Document(BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            text = "\n".join(parts)
        else:
            raise InputError("Choose a .pdf or .docx resume, or paste your resume text.")
    except InputError:
        raise
    except Exception:
        # Do not surface parser internals, paths, or document data.
        raise InputError("The file could not be read. Export a fresh PDF/DOCX or paste the text.") from None
    if len(text.strip()) < 80:
        raise InputError("Not enough readable text was found. Scanned PDFs need OCR; paste the text instead.")
    return validate_text(text, "Extracted resume", MAX_RESUME_CHARS)
