from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas

from roleready.demo import DEMO_RESUME
from roleready.errors import InputError
from roleready.extraction import MAX_FILE_BYTES, extract_resume


def pdf_bytes(text=DEMO_RESUME):
    stream = BytesIO()
    pdf = canvas.Canvas(stream)
    for i, line in enumerate(text.splitlines()):
        pdf.drawString(50, 800 - i * 15, line)
    pdf.save()
    return stream.getvalue()


def test_extract_real_text_pdf():
    assert "500 survey responses" in extract_resume(pdf_bytes(), "resume.PDF")


def test_extract_docx_paragraphs_and_tables():
    doc = Document()
    doc.add_paragraph(DEMO_RESUME)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Additional skills"
    table.cell(0, 1).text = "SQL and Python"
    out = BytesIO()
    doc.save(out)
    result = extract_resume(out.getvalue(), "resume.docx")
    assert "500 survey responses" in result
    assert "SQL and Python" in result


@pytest.mark.parametrize(
    "data,name",
    [
        (b"", "r.pdf"),
        (b"garbage", "r.pdf"),
        (b"%PDF-broken", "r.pdf"),
        (b"PKbroken", "r.docx"),
        (b"MZbad", "r.exe"),
        (b"hello", "r.docx"),
        (b"a" * (MAX_FILE_BYTES + 1), "r.pdf"),
    ],
)
def test_invalid_files_are_safe(data, name):
    with pytest.raises(InputError):
        extract_resume(data, name)


def test_image_only_or_blank_pdf():
    with pytest.raises(InputError, match="OCR"):
        extract_resume(pdf_bytes(" "), "scan.pdf")


def test_encrypted_and_long_pdf():
    writer = PdfWriter()
    writer.append(PdfReader(BytesIO(pdf_bytes())))
    writer.encrypt("secret")
    out = BytesIO()
    writer.write(out)
    with pytest.raises(InputError, match="Password"):
        extract_resume(out.getvalue(), "locked.pdf")
    writer = PdfWriter()
    for _ in range(21):
        writer.add_blank_page(width=100, height=100)
    out = BytesIO()
    writer.write(out)
    with pytest.raises(InputError, match="20 pages"):
        extract_resume(out.getvalue(), "long.pdf")


def test_docx_expansion_limit_and_non_word_zip():
    out = BytesIO()
    with ZipFile(out, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"a" * (21 * 1024 * 1024))
    with pytest.raises(InputError, match="too complex"):
        extract_resume(out.getvalue(), "bomb.docx")
    out = BytesIO()
    with ZipFile(out, "w") as archive:
        archive.writestr("other.txt", "not a document")
    with pytest.raises(InputError, match="not a Word"):
        extract_resume(out.getvalue(), "wrong.docx")


def test_macro_docx_rejected():
    out = BytesIO()
    with ZipFile(out, "w") as archive:
        archive.writestr("word/document.xml", "placeholder")
        archive.writestr("word/vbaProject.bin", "macro")
    with pytest.raises(InputError, match="Macro"):
        extract_resume(out.getvalue(), "macro.docx")
